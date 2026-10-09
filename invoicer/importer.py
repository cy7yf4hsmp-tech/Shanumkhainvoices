"""Read product names and quantities from uploaded PDF, Excel, Word or CSV files.

Tables are preferred: a header row is located by its column names (Item / Product / Particulars,
Qty / Quantity, Rate, HSN, UOM, Code ...). For PDFs without real tables, lines of text such as
"LED Bulb 9W   12 pcs   85.00" are read as a fallback and flagged for checking.
"""
import csv
import difflib
import io
import re

SUPPORTED = (".xlsx", ".xlsm", ".xls", ".csv", ".pdf", ".docx")


class ImportError_(ValueError):
    pass


# Order matters: "item code" must be a code, "unit price" a price, before the looser name/unit words match.
_HEADER_WORDS = [
    ("hsn", ("hsn", "sac")),
    ("sku", ("sku", "code", "part no", "part number", "article")),
    ("price", ("rate", "price", "mrp", "cost")),
    ("qty", ("qty", "quantity", "stock", "nos", "pcs", "count", "closing")),
    ("unit", ("uom", "unit", "per")),
    ("name", ("product", "item", "description", "particular", "name", "goods", "material", "details")),
]
_IGNORE_WORDS = ("amount", "total", "value", "gst", "tax", "disc", "s.no", "sl", "sr")


def _classify(cell):
    text = re.sub(r"\s+", " ", str(cell or "")).strip().lower()
    if not text or len(text) > 40:
        return None
    for field, words in _HEADER_WORDS:
        if any(w in text for w in words):
            # "total qty" is still a quantity, but "amount"/"taxable value" columns are not
            if field != "qty" and any(text.startswith(w) for w in _IGNORE_WORDS):
                return None
            return field
    return None


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"(?i)(rs\.?|inr|₹|,)", "", str(value)).strip()
    m = re.match(r"^-?\d+(\.\d+)?", text)
    return float(m.group(0)) if m else None


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def rows_to_items(rows):
    """Turn a table (list of rows of cells) into item dicts using its header row."""
    rows = [[_clean(c) for c in r] for r in rows if r and any(_clean(c) for c in r)]
    for h, header in enumerate(rows[:25]):
        cols = {}
        for i, cell in enumerate(header):
            field = _classify(cell)
            if field and field not in cols:
                cols[field] = i
        if "name" in cols and "qty" in cols:
            break
    else:
        return []

    items = []
    for row in rows[h + 1:]:
        def get(field):
            i = cols.get(field)
            return row[i] if i is not None and i < len(row) else ""
        name, qty = get("name"), to_number(get("qty"))
        if not name or qty is None or name.lower().startswith(("total", "grand total", "sub total", "subtotal")):
            continue
        items.append({
            "name": name, "qty": qty, "sku": get("sku"), "hsn": get("hsn"), "unit": get("unit"),
            "price": to_number(get("price")), "check": False,
        })
    return items


_LINE = re.compile(
    r"^(?:\d{1,3}[.)]?\s+)?(?P<name>[A-Za-z].*?)\s+(?P<qty>\d+(?:\.\d+)?)(?=\s|$)\s*"
    r"(?P<unit>[A-Za-z]{1,6}\b)?\s*(?P<rest>.*)$")


def text_to_items(text):
    """Fallback for documents without tables: read 'name  qty [unit] [rate ...]' lines."""
    items = []
    for line in text.splitlines():
        line = line.strip()
        m = _LINE.match(line)
        if not m or len(m.group("name")) < 2:
            continue
        rest_numbers = [to_number(x) for x in m.group("rest").split()]
        rest_numbers = [n for n in rest_numbers if n is not None]
        items.append({
            "name": m.group("name").strip(" -:|"), "qty": float(m.group("qty")), "sku": "", "hsn": "",
            "unit": m.group("unit") or "", "price": rest_numbers[0] if rest_numbers else None, "check": True,
        })
    return items


def _excel_tables(data, filename):
    if filename.endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        return [[sheet.row_values(r) for r in range(sheet.nrows)] for sheet in book.sheets()]
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [[list(r) for r in sheet.iter_rows(values_only=True)] for sheet in book.worksheets]


def _docx_content(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    return tables, "\n".join(p.text for p in doc.paragraphs)


def _pdf_content(data):
    import pdfplumber
    tables, text = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            tables.extend(page.extract_tables() or [])
            text.append(page.extract_text() or "")
    return tables, "\n".join(text)


def read_items(filename, data):
    """Return (items, notes) from an uploaded file. Raises ImportError_ with a readable message."""
    name = (filename or "").lower()
    if name.endswith(".doc"):
        raise ImportError_("Old Word .doc files can't be read. Open it in Word and use Save As → .docx, then upload again.")
    if not name.endswith(SUPPORTED):
        raise ImportError_("Please upload a PDF, Excel (.xlsx / .xls), Word (.docx) or CSV file.")
    notes = []
    try:
        if name.endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            tables, plain = [list(csv.reader(io.StringIO(text)))], ""
        elif name.endswith((".xlsx", ".xlsm", ".xls")):
            tables, plain = _excel_tables(data, name), ""
        elif name.endswith(".docx"):
            tables, plain = _docx_content(data)
        else:
            tables, plain = _pdf_content(data)
    except Exception as exc:  # corrupt or password-protected files
        raise ImportError_(f"Could not open the file: {exc}") from exc

    items = []
    for table in tables:
        items.extend(rows_to_items(table))
    if not items and plain.strip():
        items = text_to_items(plain)
        if items:
            notes.append("No table with Item and Qty columns was found, so the text was read line by line. "
                         "Please check every row before applying.")
    if not items:
        if name.endswith(".pdf") and not plain.strip():
            raise ImportError_("This PDF has no readable text (it looks like a scanned image). "
                               "Please upload the Excel/Word version or a PDF created from software.")
        raise ImportError_("No items found. The file needs a table with an item/product column and a quantity column.")
    return items, notes


def _norm(text):
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def match_product(item, products):
    """Find the product an imported row refers to. Returns (product_id or None, how)."""
    sku = _norm(item.get("sku"))
    if sku:
        for p in products:
            if _norm(p["sku"]) == sku:
                return p["id"], "code"
    name = _norm(item["name"])
    for p in products:
        if _norm(p["name"]) == name or (p["sku"] and _norm(p["sku"]) == name):
            return p["id"], "name"
    names = {_norm(p["name"]): p["id"] for p in products}
    close = difflib.get_close_matches(name, list(names), n=1, cutoff=0.8)
    if close:
        return names[close[0]], "similar"
    return None, ""
