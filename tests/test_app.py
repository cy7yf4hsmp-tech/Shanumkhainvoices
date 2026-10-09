import pytest

from invoicer import create_app
from invoicer.db import get_db
from invoicer.utils import amount_in_words, inr


@pytest.fixture
def app(tmp_path):
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "test.sqlite3"), "SECRET_KEY": "test",
                      "CSRF_ENABLED": False})
    with app.app_context():
        make_user(get_db(), "admin", "admin")
    return app


def make_user(db, username, role, perms=None, password="password123"):
    from werkzeug.security import generate_password_hash
    uid = db.execute("INSERT INTO users (username, full_name, password_hash, role) VALUES (?, ?, ?, ?)",
                     (username, username.title(), generate_password_hash(password), role)).lastrowid
    for section, level in (perms or {}).items():
        db.execute("INSERT INTO user_permissions (user_id, section, level) VALUES (?, ?, ?)", (uid, section, level))
    db.commit()
    return uid


def login(client, username, password="password123"):
    return client.post("/login", data={"username": username, "password": password})


@pytest.fixture
def client(app):
    """A browser logged in as the administrator."""
    c = app.test_client()
    assert login(c, "admin").status_code == 302
    return c


def add_product(client, **kw):
    data = {"name": "Widget", "sku": "W1", "hsn": "8471", "unit": "pcs", "price": "100", "cost": "60",
            "gst_rate": "18", "stock": "10", "reorder_level": "2"} | kw
    return client.post("/products/new", data=data, follow_redirects=True)


def stock_of(app, pid=1):
    with app.app_context():
        return get_db().execute("SELECT stock FROM products WHERE id = ?", (pid,)).fetchone()[0]


def invoice_data(**kw):
    return {"customer_name": "Ravi Traders", "tax_type": "intra", "date": "2026-10-09",
            "product_id": ["1"], "description": ["Widget"], "hsn": ["8471"], "unit": ["pcs"],
            "qty": ["3"], "price": ["100"], "discount_pct": ["0"], "gst_rate": ["18"]} | kw


def test_pages_render(client):
    add_product(client)
    for url in ["/", "/products", "/products/1", "/stock", "/stock?view=out&period=today", "/stock?view=all&period=fy",
                "/stock?period=custom&from=2026-01-01&to=2026-12-31", "/customers", "/customers/new",
                "/invoices", "/invoices/new", "/settings", "/stock/export.csv", "/payments", "/products/import"]:
        assert client.get(url).status_code == 200, url


def test_opening_stock_recorded(client, app):
    add_product(client)
    assert stock_of(app) == 10
    with app.app_context():
        kinds = [r[0] for r in get_db().execute("SELECT kind FROM stock_movements")]
    assert kinds == ["opening"]


def test_invoice_deducts_stock_and_computes_gst(client, app):
    add_product(client)
    resp = client.post("/invoices/new", data=invoice_data(), follow_redirects=True)
    assert resp.status_code == 200
    assert b"SCP/2026-27/10/001" in resp.data
    assert stock_of(app) == 7
    with app.app_context():
        inv = get_db().execute("SELECT * FROM invoices").fetchone()
    assert inv["subtotal"] == 300
    assert inv["tax_total"] == 54
    assert inv["total"] == 354
    assert inv["status"] == "unpaid"


def test_insufficient_stock_rejected_and_nothing_saved(client, app):
    add_product(client)
    resp = client.post("/invoices/new", data=invoice_data(qty=["50"]))
    assert resp.status_code == 400
    assert b"Not enough stock" in resp.data
    assert stock_of(app) == 10
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM invoices").fetchone()[0] == 0


def test_cancel_restores_stock(client, app):
    add_product(client)
    client.post("/invoices/new", data=invoice_data())
    client.post("/invoices/1/cancel")
    assert stock_of(app) == 10
    client.post("/invoices/1/cancel")  # idempotent
    assert stock_of(app) == 10


def test_payments_update_status(client, app):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(amount_paid="100"))
    with app.app_context():
        assert get_db().execute("SELECT status FROM invoices").fetchone()[0] == "partial"
    client.post("/invoices/1/pay", data={"amount": "254"})
    with app.app_context():
        assert get_db().execute("SELECT status, amount_paid FROM invoices").fetchone()[:] == ("paid", 354)


def test_stock_in_out_and_count(client, app):
    add_product(client)
    client.post("/products/1/stock", data={"action": "in", "qty": "5"})
    assert stock_of(app) == 15
    client.post("/products/1/stock", data={"action": "out", "qty": "3"})
    assert stock_of(app) == 12
    client.post("/products/1/stock", data={"action": "count", "qty": "8"})
    assert stock_of(app) == 8
    resp = client.post("/products/1/stock", data={"action": "out", "qty": "100"}, follow_redirects=True)
    assert b"Not enough stock" in resp.data
    assert stock_of(app) == 8


def test_custom_line_without_product_and_discount(client, app):
    data = invoice_data(product_id=[""], description=["Installation"], qty=["1"], price=["1000"],
                        discount_pct=["10"], gst_rate=["18"], tax_type="inter")
    client.post("/invoices/new", data=data)
    with app.app_context():
        inv = get_db().execute("SELECT * FROM invoices").fetchone()
    assert (inv["discount"], inv["tax_total"], inv["total"]) == (100, 162, 1062)


def test_invoice_numbers_increment(client):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"]))
    resp = client.post("/invoices/new", data=invoice_data(qty=["1"]), follow_redirects=True)
    assert b"SCP/2026-27/10/002" in resp.data


def test_formatting_helpers():
    assert inr(1234567.5) == "12,34,567.50"
    assert inr(999) == "999.00"
    assert amount_in_words(354) == "Rupees Three Hundred Fifty Four Only"
    assert amount_in_words(1250000.25) == "Rupees Twelve Lakh Fifty Thousand and Twenty Five Paise Only"


def test_new_pages_render(client):
    for url in ["/dispatch", "/expenses", "/loans", "/loans/new"]:
        assert client.get(url).status_code == 200, url


def test_invoice_list_shows_open_by_default(client):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"], customer_name="Open Co"))
    client.post("/invoices/new", data=invoice_data(qty=["1"], customer_name="Paid Co", amount_paid="118"))
    page = client.get("/invoices").data
    assert b"Open Co" in page and b"Paid Co" not in page
    page = client.get("/invoices?status=all").data
    assert b"Open Co" in page and b"Paid Co" in page


def test_ship_to_address(client, app):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"], customer_address="Bill St", ship_same="1",
                                                   ship_to_address="ignored"))
    client.post("/invoices/new", data=invoice_data(qty=["1"], customer_address="Bill St",
                                                   ship_to_address="Godown, Ring Road"))
    with app.app_context():
        rows = [r[0] for r in get_db().execute("SELECT ship_to_address FROM invoices ORDER BY id")]
    assert rows == ["Bill St", "Godown, Ring Road"]
    assert b"Godown, Ring Road" in client.get("/invoices/2").data


def test_item_name_comes_from_product(client, app):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(description=[""]))
    with app.app_context():
        assert get_db().execute("SELECT description FROM invoice_items").fetchone()[0] == "Widget"


def test_printed_invoices_appear_in_dispatch(client, app):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"]))
    assert b"SCP/2026-27/10/001" not in client.get("/dispatch").data
    client.get("/invoices/1/print")
    assert b"SCP/2026-27/10/001" in client.get("/dispatch").data
    client.post("/dispatch/1", data={"dispatch_status": "dispatched", "transport_type": "Auto"})
    with app.app_context():
        row = get_db().execute("SELECT dispatch_status, transport_type, dispatch_date FROM invoices").fetchone()
    assert row["dispatch_status"] == "dispatched" and row["transport_type"] == "Auto" and row["dispatch_date"]
    assert b"SCP/2026-27/10/001" not in client.get("/dispatch?status=not_dispatched").data


def test_expenses(client, app):
    client.post("/expenses", data={"date": "2026-10-05", "category": "Auto", "amount": "150"})
    client.post("/expenses", data={"date": "2026-10-06", "category": "Transportation", "amount": "2000"})
    client.post("/expenses", data={"date": "2026-10-06", "category": "Other", "amount": "0"})  # rejected
    page = client.get("/expenses?month=2026-10").data
    assert b"2,150.00" in page
    client.post("/expenses/1/delete")
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM expenses").fetchone()[0] == 1


def test_loans_with_shares_and_repayments(client, app):
    resp = client.post("/loans/new", data={
        "date": "2026-10-01", "lender_type": "bank", "lender_name": "SBI", "amount": "500000",
        "interest_rate": "10.5", "share_person": ["Ravi", "Suresh", ""], "share_amount": ["300000", "200000", ""],
    }, follow_redirects=True)
    assert b"SBI" in resp.data and b"60.0%" in resp.data
    client.post("/loans/1/repay", data={"date": "2026-10-09", "amount": "25000"})
    page = client.get("/loans").data
    assert b"4,75,000.00" in page and b"Ravi" in page
    # shares that add up to more than the loan are refused
    resp = client.post("/loans/new", data={"lender_type": "person", "lender_name": "Ramesh", "amount": "1000",
                                           "share_person": ["A"], "share_amount": ["5000"]})
    assert b"more than the loan amount" in resp.data
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM loans").fetchone()[0] == 1


def test_old_database_is_upgraded(tmp_path):
    import sqlite3
    path = tmp_path / "old.sqlite3"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE invoices (id INTEGER PRIMARY KEY, number TEXT, date TEXT, customer_name TEXT)")
    con.commit()
    con.close()
    app = create_app({"TESTING": True, "DATABASE": str(path)})
    with app.app_context():
        cols = {r["name"] for r in get_db().execute("PRAGMA table_info(invoices)")}
    assert {"ship_to_address", "printed_at", "dispatch_status", "transport_type"} <= cols


# ---------------------------------------------------------------- invoice numbering, due days, payments

def test_invoice_number_format_and_financial_year(client, app):
    add_product(client, stock="100")
    for d in ["2027-03-31", "2027-04-01", "2027-04-15", "2027-05-02"]:
        client.post("/invoices/new", data=invoice_data(qty=["1"], date=d))
    with app.app_context():
        numbers = [r[0] for r in get_db().execute("SELECT number FROM invoices ORDER BY id")]
    assert numbers == ["SCP/2026-27/03/001", "SCP/2027-28/04/001", "SCP/2027-28/04/002", "SCP/2027-28/05/003"]


def test_next_number_can_be_set_in_settings(client, app):
    from invoicer import services
    add_product(client)
    with app.app_context():
        settings = dict(get_db().execute("SELECT key, value FROM settings").fetchall())
    settings.update({"next_invoice_no": "46", "default_due_days": "15"})
    client.post("/settings", data=settings)
    with app.app_context():
        fy = services.financial_year(__import__("datetime").date.today())
        assert services.next_invoice_seq(get_db(), fy) == 46


def test_due_days_sets_due_date(client, app):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"], date="2026-10-09", due_days="30"))
    with app.app_context():
        row = get_db().execute("SELECT due_days, due_date FROM invoices").fetchone()
    assert (row["due_days"], row["due_date"]) == (30, "2026-11-08")


def test_payments_page(client):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"], customer_name="Late Payer", due_days="0",
                                                   date="2026-01-01"))
    client.post("/invoices/new", data=invoice_data(qty=["1"], customer_name="Good Payer", amount_paid="118"))
    page = client.get("/payments").data.decode()
    assert "Late Payer" in page and "Overdue by" in page and "Good Payer" not in page
    page = client.get("/payments?status=received").data.decode()
    assert "Good Payer" in page and "Payment received" in page
    client.post("/invoices/1/pay", data={"amount": "118", "next": "/payments"})
    assert "Late Payer" not in client.get("/payments").data.decode()


def test_customers_have_city(client, app):
    client.post("/customers/new", data={"name": "Sri Sai Traders", "address": "Main Road", "city": "Guntur",
                                        "email": "sai@example.com", "phone": "90000 11111"})
    page = client.get("/customers").data.decode()
    for text in ["Sri Sai Traders", "Main Road", "Guntur", "sai@example.com", "90000 11111", "Mail ID", "City"]:
        assert text in page
    assert "Sri Sai Traders" in client.get("/customers?q=guntur").data.decode()


# ---------------------------------------------------------------- stock summary & dashboard stock out

def test_stock_summary_and_dashboard_stock_out(client, app):
    from datetime import date
    add_product(client, stock="10")
    today = date.today().isoformat()
    client.post("/invoices/new", data=invoice_data(qty=["3"], date=today))
    client.post("/products/1/stock", data={"action": "out", "qty": "1", "note": "damaged"})
    client.post("/products/1/stock", data={"action": "in", "qty": "5"})
    page = client.get("/stock?period=today").data.decode()
    assert "−4" in page and "+15" in page  # out: 3 sold + 1 damaged; in: 10 opening + 5 purchased (all today)
    out = client.get("/stock?view=out&period=today").data.decode()
    assert "damaged" in out and "Ravi Traders" in out
    dash = client.get("/").data.decode()
    assert "Stock out" in dash and "This week" in dash and "₹400.00" in dash  # 4 units x ₹100


# ---------------------------------------------------------------- uploading stock files

def _xlsx_bytes(rows):
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Supplier: ABC Distributors"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(rows):
    import io
    from docx import Document
    doc = Document()
    doc.add_paragraph("Purchase bill")
    table = doc.add_table(rows=0, cols=len(rows[0]))
    for r in rows:
        cells = table.add_row().cells
        for c, v in zip(cells, r):
            c.text = str(v)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(lines):
    """A minimal one-page PDF with plain text lines (no table)."""
    text = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({l}) Tj T*" for l in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> "
            "/Contents 5 0 R >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            f"<< /Length {len(text)} >>\nstream\n{text}\nendstream"]
    out, offsets = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
    return out.encode("latin-1")


def test_importer_reads_all_file_types():
    from invoicer.importer import read_items
    header = ["S.No", "Item Code", "Description of Goods", "HSN", "Qty", "UOM", "Rate", "Amount"]
    rows = [header, [1, "W1", "Widget", "8471", 12, "pcs", 100, 1200], [2, "", "Gadget Pro", "", "5", "box", "250", "1250"],
            ["", "", "Total", "", 17, "", "", 2450]]
    for name, data in [("bill.xlsx", _xlsx_bytes(rows)), ("bill.docx", _docx_bytes(rows)),
                       ("bill.csv", "\n".join(",".join(str(c) for c in r) for r in rows).encode())]:
        items, notes = read_items(name, data)
        assert [(i["name"], i["qty"], i["sku"]) for i in items] == [("Widget", 12, "W1"), ("Gadget Pro", 5, "")], name
        assert items[0]["price"] == 100 and items[1]["unit"] == "box"
    items, notes = read_items("bill.pdf", _pdf_bytes(["ABC Distributors", "1 Widget 12 pcs 100.00",
                                                      "2 LED Bulb 9W 30 pcs 85.00"]))
    assert [(i["name"], i["qty"]) for i in items] == [("Widget", 12), ("LED Bulb 9W", 30)]
    assert notes  # line-by-line reading is flagged for checking


def test_importer_rejects_bad_files():
    import pytest as _pytest
    from invoicer.importer import ImportError_, read_items
    for name, data, msg in [("old.doc", b"x", "Save As"), ("pic.png", b"x", "Please upload"),
                            ("empty.csv", b"just,some\nwords,here", "No items found")]:
        with _pytest.raises(ImportError_, match=msg):
            read_items(name, data)


def test_upload_preview_and_apply(client, app):
    import io
    add_product(client, stock="10")  # Widget, code W1
    rows = [["Item Code", "Particulars", "Qty", "Rate"], ["W1", "Widget (blue)", 12, 100], ["", "Brand New Thing", 4, 55]]
    resp = client.post("/products/import", data={"file": (io.BytesIO(_xlsx_bytes(rows)), "bill.xlsx"), "mode": "add"},
                       content_type="multipart/form-data")
    page = resp.data.decode()
    assert resp.status_code == 200 and "by code" in page and "not in products" in page
    assert stock_of(app) == 10  # preview only, nothing changed yet
    client.post("/products/import/apply", data={
        "mode": "add", "reference": "bill.xlsx", "target": ["1", "new"], "name": ["Widget (blue)", "Brand New Thing"],
        "sku": ["W1", ""], "hsn": ["", ""], "unit": ["", "pcs"], "price": ["100", "55"], "qty": ["12", "4"]})
    assert stock_of(app) == 22
    with app.app_context():
        new = get_db().execute("SELECT * FROM products WHERE name = 'Brand New Thing'").fetchone()
    assert new["stock"] == 4 and new["price"] == 55
    # "replace stock" mode sets the count
    client.post("/products/import/apply", data={"mode": "count", "target": ["1"], "name": ["Widget"], "sku": [""],
                                                "hsn": [""], "unit": [""], "price": [""], "qty": ["7"]})
    assert stock_of(app) == 7


# ---------------------------------------------------------------- logins, roles and access

def test_must_log_in(app):
    c = app.test_client()
    resp = c.get("/invoices")
    assert resp.status_code == 302 and "/login" in resp.headers["Location"]
    assert c.post("/expenses", data={"amount": "5"}).status_code == 302  # no saving without login
    assert login(c, "admin", "wrong-password").status_code == 401
    assert c.get("/invoices").status_code == 302


def test_first_run_setup_creates_admin(tmp_path):
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "new.sqlite3"), "CSRF_ENABLED": False})
    c = app.test_client()
    assert "/setup" in c.get("/").headers["Location"]
    resp = c.post("/setup", data={"full_name": "Owner", "username": "owner", "password": "longpassword",
                                  "confirm_password": "longpassword"})
    assert resp.status_code == 302
    assert c.get("/users").status_code == 200  # logged in as the new admin
    # setup can't be used again once a user exists
    assert "/login" in app.test_client().get("/setup").headers["Location"]


def test_lockout_after_repeated_wrong_passwords(app):
    c = app.test_client()
    for _ in range(5):
        login(c, "admin", "nope")
    resp = login(c, "admin")  # right password, but locked
    assert resp.status_code == 401 and b"Too many wrong passwords" in resp.data
    with app.app_context():
        get_db().execute("UPDATE users SET locked_until = '2000-01-01 00:00:00'")
        get_db().commit()
    assert login(c, "admin").status_code == 302


def test_staff_only_sees_allowed_sections(app):
    with app.app_context():
        make_user(get_db(), "ravi", "staff", {"invoices": "edit", "customers": "view"})
    c = app.test_client()
    login(c, "ravi")
    home = c.get("/", follow_redirects=True)
    assert b"Invoices" in home.data and b"Loans" not in home.data and b"Users &amp; access" not in home.data
    assert c.get("/invoices").status_code == 200
    assert c.get("/invoices/new").status_code == 200
    assert c.get("/loans").status_code == 403
    assert c.get("/payments").status_code == 403
    assert c.get("/users").status_code == 403
    assert c.post("/expenses", data={"amount": "100", "category": "Auto"}).status_code == 403
    # view-only on customers: can look, can't add or change
    assert c.get("/customers").status_code == 200
    assert b"+ Add customer" not in c.get("/customers").data
    assert c.get("/customers/new").status_code == 403
    assert c.post("/customers/new", data={"name": "X"}).status_code == 403


def test_investor_view_only(app, client):
    client.post("/loans/new", data={"lender_type": "bank", "lender_name": "SBI", "amount": "1000"})
    with app.app_context():
        make_user(get_db(), "investor1", "investor", {"loans": "view", "dashboard": "view"})
    c = app.test_client()
    login(c, "investor1")
    assert c.get("/loans").status_code == 200 and b"SBI" in c.get("/loans").data
    assert b"Record a loan" not in c.get("/loans").data
    assert c.get("/loans/new").status_code == 403
    assert c.post("/loans/1/repay", data={"amount": "10"}).status_code == 403
    assert c.post("/loans/1/delete").status_code == 403


def test_no_dashboard_access_goes_to_first_allowed_page(app):
    with app.app_context():
        make_user(get_db(), "packer", "staff", {"dispatch": "edit"})
    c = app.test_client()
    resp = login(c, "packer")
    assert resp.headers["Location"].endswith("/dispatch")
    assert c.get("/").headers["Location"].endswith("/dispatch")


def test_admin_creates_user_and_sets_access(app, client):
    resp = client.post("/users/new", data={
        "full_name": "Suresh", "username": "suresh", "role": "partner", "password": "secret-pass",
        "confirm_password": "secret-pass", "active": "1", "perm_invoices": "edit", "perm_loans": "view",
        "perm_settings": "none"})
    assert resp.status_code == 302
    c = app.test_client()
    assert login(c, "suresh", "secret-pass").status_code == 302
    assert c.get("/loans").status_code == 200 and c.get("/loans/new").status_code == 403
    assert c.get("/settings").status_code == 403
    # admin changes access; it applies on the very next page
    with app.app_context():
        uid = get_db().execute("SELECT id FROM users WHERE username = 'suresh'").fetchone()[0]
    client.post(f"/users/{uid}/edit", data={"full_name": "Suresh", "role": "partner", "active": "1",
                                            "perm_loans": "edit"})
    assert c.get("/loans/new").status_code == 200
    assert c.get("/invoices").status_code == 403
    # deactivating logs them out
    client.post(f"/users/{uid}/edit", data={"full_name": "Suresh", "role": "partner"})
    assert c.get("/loans").status_code == 302


def test_admin_can_reset_password_and_username_is_unique(app, client):
    client.post("/users/new", data={"full_name": "A", "username": "anil", "role": "staff",
                                    "password": "first-pass", "confirm_password": "first-pass", "active": "1"})
    dup = client.post("/users/new", data={"full_name": "B", "username": "ANIL", "role": "staff",
                                          "password": "other-pass", "confirm_password": "other-pass"})
    assert dup.status_code == 400 and b"already taken" in dup.data
    with app.app_context():
        uid = get_db().execute("SELECT id FROM users WHERE username = 'anil'").fetchone()[0]
    client.post(f"/users/{uid}/edit", data={"full_name": "A", "role": "staff", "active": "1",
                                            "password": "new-pass-1", "confirm_password": "new-pass-1"})
    c = app.test_client()
    assert login(c, "anil", "first-pass").status_code == 401
    assert login(c, "anil", "new-pass-1").status_code == 302


def test_last_admin_cannot_be_removed(app, client):
    resp = client.post("/users/1/edit", data={"full_name": "Admin", "role": "staff", "active": "1"})
    assert resp.status_code == 400 and b"at least one active administrator" in resp.data
    with app.app_context():
        assert get_db().execute("SELECT role FROM users WHERE id = 1").fetchone()[0] == "admin"


def test_change_own_password(app, client):
    bad = client.post("/account", data={"current_password": "wrong", "new_password": "abcdefgh1",
                                        "confirm_password": "abcdefgh1"}, follow_redirects=True)
    assert b"not correct" in bad.data
    client.post("/account", data={"current_password": "password123", "new_password": "abcdefgh1",
                                  "confirm_password": "abcdefgh1"})
    c = app.test_client()
    assert login(c, "admin", "abcdefgh1").status_code == 302


def test_changes_record_who_made_them(app, client):
    add_product(client)
    client.post("/invoices/new", data=invoice_data(qty=["1"]))
    client.post("/expenses", data={"category": "Auto", "amount": "50"})
    with app.app_context():
        db = get_db()
        assert db.execute("SELECT created_by FROM invoices").fetchone()[0] == "Admin"
        assert db.execute("SELECT user FROM stock_movements WHERE kind = 'sale'").fetchone()[0] == "Admin"
        assert db.execute("SELECT created_by FROM expenses").fetchone()[0] == "Admin"


def test_csrf_protection(tmp_path):
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "c.sqlite3"), "SECRET_KEY": "t"})
    with app.app_context():
        make_user(get_db(), "admin", "admin")
    c = app.test_client()
    assert c.post("/login", data={"username": "admin", "password": "password123"}).status_code == 400
    import re
    token = re.search(rb'name="_csrf" value="([^"]+)"', c.get("/login").data).group(1).decode()
    assert c.post("/login", data={"username": "admin", "password": "password123", "_csrf": token}).status_code == 302
    assert c.post("/expenses", data={"category": "Auto", "amount": "5"}).status_code == 400  # forged form
    # logging in starts a fresh session with a new token, which the next page carries
    token = re.search(rb'name="_csrf" value="([^"]+)"', c.get("/expenses").data).group(1).decode()
    assert c.post("/expenses", data={"category": "Auto", "amount": "5", "_csrf": token}).status_code == 302
